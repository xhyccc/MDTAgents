"""
context_extractor.py — Deterministic pre-processing of complex files.

For each PDF / Word / Excel file found in a case, creates a context
sub-folder under ``.mdt_workspace/context/`` that contains:

  <stem>_<ext>/
  ├── <stem>.txt           full extracted text
  ├── page_001.png         page-1 screenshot  (PDF only, requires PyMuPDF)
  ├── page_002.png         …
  └── image_001.png        embedded raster image (PDF / DOCX)

Runs *before* any AI work so every agent's workspace is pre-populated.
All operations are idempotent: existing output files are never overwritten.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable, List, Optional


COMPLEX_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx"})


class ContextExtractor:
    """
    Extract text, page screenshots, and embedded images from complex files.

    Parameters
    ----------
    context_dir:
        Root context directory, typically ``.mdt_workspace/context/``.
    """

    def __init__(self, context_dir: Path) -> None:
        self.context_dir = Path(context_dir)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def extract_all(
        self,
        file_paths: List[Path],
        progress_cb: Optional[Callable[[str], None]] = None,
    ) -> None:
        """Extract context for every complex file in *file_paths*.

        Parameters
        ----------
        file_paths:
            Absolute paths to files to process.
        progress_cb:
            Optional callable(message) for progress reporting (e.g. Streamlit st.write).
        """
        complex_files = [p for p in file_paths if p.suffix.lower() in COMPLEX_SUFFIXES]
        for p in complex_files:
            out_dir = self.file_context_dir(p)
            if out_dir.exists() and any(out_dir.iterdir()):
                if progress_cb:
                    progress_cb(f"⏭ `{p.name}` — already extracted")
                continue
            if progress_cb:
                progress_cb(f"⚙ `{p.name}` …")
            try:
                self.extract_file(p)
                if progress_cb:
                    progress_cb(f"✅ `{p.name}` → `{out_dir.name}/`")
            except Exception as exc:
                if progress_cb:
                    progress_cb(f"⚠ `{p.name}` extraction error: {exc}")

    def file_context_dir(self, file_path: Path) -> Path:
        """Return context subfolder for *file_path*.

        E.g. ``报告.pdf`` → ``context/报告_pdf/``
        """
        stem = file_path.stem
        ext = file_path.suffix.lstrip(".").lower() or "file"
        return self.context_dir / f"{stem}_{ext}"

    def extract_file(self, file_path: Path) -> Path:
        """Extract one file; create and return its context directory."""
        out_dir = self.file_context_dir(file_path)
        out_dir.mkdir(parents=True, exist_ok=True)
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            self._extract_pdf(file_path, out_dir)
        elif suffix == ".docx":
            self._extract_docx(file_path, out_dir)
        elif suffix == ".xlsx":
            self._extract_xlsx(file_path, out_dir)
        return out_dir

    # ------------------------------------------------------------------
    # Per-format extractors
    # ------------------------------------------------------------------

    def _extract_pdf(self, path: Path, out_dir: Path) -> None:
        # ── Full text ─────────────────────────────────────────────────
        txt_path = out_dir / (path.stem + ".txt")
        if not txt_path.exists():
            txt_path.write_text(self._pdf_text(path), encoding="utf-8")

        # ── Page screenshots + embedded images via PyMuPDF ────────────
        try:
            import fitz  # PyMuPDF  # type: ignore
        except ImportError:
            return

        doc = fitz.open(str(path))
        img_counter = 1
        seen_xrefs: set = set()

        for page_no, page in enumerate(doc, start=1):
            # Per-page screenshot
            png_path = out_dir / f"page_{page_no:03d}.png"
            if not png_path.exists():
                mat = fitz.Matrix(2.0, 2.0)  # 2× zoom ≈ 150 DPI
                pix = page.get_pixmap(matrix=mat)
                pix.save(str(png_path))

            # Embedded raster images on this page
            for img_info in page.get_images(full=True):
                xref = img_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                base_img = doc.extract_image(xref)
                img_bytes = base_img.get("image", b"")
                img_ext = base_img.get("ext", "png")
                img_path = out_dir / f"image_{img_counter:03d}.{img_ext}"
                if not img_path.exists() and img_bytes:
                    img_path.write_bytes(img_bytes)
                img_counter += 1

        doc.close()

    def _extract_docx(self, path: Path, out_dir: Path) -> None:
        # ── Full text ─────────────────────────────────────────────────
        txt_path = out_dir / (path.stem + ".txt")
        if not txt_path.exists():
            txt_path.write_text(self._docx_text(path), encoding="utf-8")

        # ── Embedded images (from the .docx zip) ──────────────────────
        import zipfile

        img_suffixes = frozenset({".png", ".jpg", ".jpeg", ".gif", ".webp", ".tiff"})
        img_counter = 1
        try:
            with zipfile.ZipFile(str(path), "r") as zf:
                for entry in sorted(zf.namelist()):
                    if not entry.startswith("word/media/"):
                        continue
                    img_suffix = Path(entry).suffix.lower()
                    if img_suffix not in img_suffixes:
                        continue
                    img_path = out_dir / f"image_{img_counter:03d}{img_suffix}"
                    if not img_path.exists():
                        img_path.write_bytes(zf.read(entry))
                    img_counter += 1
        except Exception:
            pass

    def _extract_xlsx(self, path: Path, out_dir: Path) -> None:
        txt_path = out_dir / (path.stem + ".txt")
        if not txt_path.exists():
            txt_path.write_text(self._xlsx_text(path), encoding="utf-8")

    # ------------------------------------------------------------------
    # Text extraction helpers (self-contained, no dependency on scanner)
    # ------------------------------------------------------------------

    @staticmethod
    def _pdf_text(path: Path) -> str:
        try:
            import pdfplumber  # type: ignore
            with pdfplumber.open(path) as pdf:
                parts = [page.extract_text() or "" for page in pdf.pages]
            return "\n\n".join(p for p in parts if p.strip())
        except ImportError:
            return "[PDF extraction requires pdfplumber: pip install pdfplumber]"
        except Exception as exc:
            return f"[PDF read error: {exc}]"

    @staticmethod
    def _docx_text(path: Path) -> str:
        try:
            from docx import Document  # type: ignore
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return "[DOCX extraction requires python-docx: pip install python-docx]"
        except Exception as exc:
            return f"[DOCX read error: {exc}]"

    @staticmethod
    def _xlsx_text(path: Path) -> str:
        try:
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows: List[str] = []
            for ws in wb.worksheets:
                rows.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(rows)
        except ImportError:
            return "[XLSX extraction requires openpyxl: pip install openpyxl]"
        except Exception as exc:
            return f"[XLSX read error: {exc}]"
