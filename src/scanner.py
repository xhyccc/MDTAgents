"""
scanner.py — Pure file scanner with zero business logic.

Recursively scans a case directory, extracts text previews, and generates
a manifest (00_manifest.json). No medical keywords, no regex-based type
detection. All type classification is delegated to the Coordinator Agent.
"""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional


# ---------------------------------------------------------------------------
# Data models
# ---------------------------------------------------------------------------

PREVIEW_CHARS = 800  # max characters per file preview
WORKSPACE_DIR_NAME = ".mdt_workspace"


@dataclass
class FileEntry:
    path: str          # relative path from case_dir
    size: int          # bytes
    mime_type: str
    preview: str       # first PREVIEW_CHARS characters of text content
    checksum: str      # MD5 hex digest


@dataclass
class Manifest:
    case_id: str
    files: List[FileEntry] = field(default_factory=list)
    total_files: int = 0
    timestamp: str = ""

    def to_dict(self) -> dict:
        return {
            "case_id": self.case_id,
            "files": [asdict(f) for f in self.files],
            "total_files": self.total_files,
            "timestamp": self.timestamp,
        }

    def to_json(self, indent: int = 2) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=indent)


# ---------------------------------------------------------------------------
# Text extractors (one per MIME / extension)
# ---------------------------------------------------------------------------

def _extract_text_plain(path: Path) -> str:
    """Read plain-text files (md, txt, csv, json, etc.)."""
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""


def _extract_text_pdf(path: Path) -> str:
    """Extract text from a PDF (text layer only, no OCR)."""
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(path) as pdf:
            parts = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
                    if sum(len(p) for p in parts) >= PREVIEW_CHARS * 2:
                        break
            return "\n".join(parts)
    except ImportError:
        return "[PDF extraction requires pdfplumber: pip install pdfplumber]"
    except Exception as exc:
        return f"[PDF read error: {exc}]"


def _extract_text_docx(path: Path) -> str:
    """Extract text from a .docx file."""
    try:
        from docx import Document  # type: ignore
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return "[DOCX extraction requires python-docx: pip install python-docx]"
    except Exception as exc:
        return f"[DOCX read error: {exc}]"


def _extract_text_xlsx(path: Path) -> str:
    """Convert an .xlsx file to a plain-text table representation."""
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows: List[str] = []
        for ws in wb.worksheets:
            rows.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join("" if v is None else str(v) for v in row))
                if sum(len(r) for r in rows) >= PREVIEW_CHARS * 2:
                    break
        return "\n".join(rows)
    except ImportError:
        return "[XLSX extraction requires openpyxl: pip install openpyxl]"
    except Exception as exc:
        return f"[XLSX read error: {exc}]"


def _extract_text_html(path: Path) -> str:
    """Convert HTML to plain text via html2text."""
    try:
        import html2text  # type: ignore
        h = html2text.HTML2Text()
        h.ignore_links = False
        h.ignore_images = True
        h.body_width = 0
        raw = path.read_text(encoding="utf-8", errors="replace")
        return h.handle(raw)
    except ImportError:
        try:
            raw = path.read_text(encoding="utf-8", errors="replace")
            import re
            return re.sub(r"<[^>]+>", "", raw)
        except Exception as exc:
            return f"[HTML read error: {exc}]"
    except Exception as exc:
        return f"[HTML read error: {exc}]"


def _extract_text_image(path: Path) -> str:
    """Return a placeholder for image files — actual content passed via --file."""
    return f"[Image file: {path.name} — will be passed as visual attachment]"


# Image suffixes handled natively as --file attachments to the model
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".webp", ".gif"})

# Map of file suffixes → extractor functions
_SUFFIX_EXTRACTORS = {
    ".md": _extract_text_plain,
    ".txt": _extract_text_plain,
    ".json": _extract_text_plain,
    ".csv": _extract_text_plain,
    ".pdf": _extract_text_pdf,
    ".docx": _extract_text_docx,
    ".xlsx": _extract_text_xlsx,
    ".html": _extract_text_html,
    ".htm": _extract_text_html,
    ".png": _extract_text_image,
    ".jpg": _extract_text_image,
    ".jpeg": _extract_text_image,
    ".webp": _extract_text_image,
    ".gif": _extract_text_image,
}


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    extractor = _SUFFIX_EXTRACTORS.get(suffix)
    if extractor is None:
        return f"[Unsupported format: {suffix}]"
    return extractor(path)


def _pdf_to_images(pdf_path: Path) -> List[Path]:
    """Render each page of a scanned PDF to a PNG image beside the original file."""
    try:
        import fitz  # type: ignore  # PyMuPDF
        doc = fitz.open(str(pdf_path))
        stem = pdf_path.stem
        out_dir = pdf_path.parent
        image_paths: List[Path] = []
        for i, page in enumerate(doc):
            mat = fitz.Matrix(2.0, 2.0)  # 2× zoom → ~150 DPI
            pix = page.get_pixmap(matrix=mat)
            out_path = out_dir / f"{stem}_page_{i + 1:03d}.png"
            pix.save(str(out_path))
            image_paths.append(out_path)
        doc.close()
        return image_paths
    except ImportError:
        return []
    except Exception:
        return []


def _checksum(path: Path) -> str:
    h = hashlib.md5()
    try:
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(65536), b""):
                h.update(chunk)
    except OSError:
        pass
    return h.hexdigest()


def _mime_type(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    return mime or "application/octet-stream"


# ---------------------------------------------------------------------------
# Scanner
# ---------------------------------------------------------------------------

class Scanner:
    """
    Recursively scans a case directory and builds a Manifest.

    Constraints:
    - No medical keywords anywhere in this file.
    - No regex-based filename matching for type detection.
    - The .mdt_workspace directory is always excluded from scanning.
    """

    SUPPORTED_SUFFIXES = frozenset(_SUFFIX_EXTRACTORS.keys())

    def scan(self, case_dir: Path) -> Manifest:
        """Scan *case_dir* and return a populated Manifest.

        For scanned (image-only) PDFs, renders each page as a PNG and adds
        those images to the manifest so they can be passed as visual attachments.
        """
        case_dir = Path(case_dir).resolve()
        manifest = Manifest(
            case_id=case_dir.name,
            timestamp=datetime.now(timezone.utc).isoformat(),
        )

        for root, dirs, files in os.walk(case_dir):
            # Always skip the workspace directory
            dirs[:] = [d for d in dirs if d != WORKSPACE_DIR_NAME]

            root_path = Path(root)
            for filename in sorted(files):
                file_path = root_path / filename
                suffix = file_path.suffix.lower()
                if suffix not in self.SUPPORTED_SUFFIXES:
                    continue

                rel_path = str(file_path.relative_to(case_dir))
                size = file_path.stat().st_size
                mime = _mime_type(file_path)
                raw_text = _extract_text(file_path)

                # For scanned PDFs (no extractable text), render pages as images
                if suffix == ".pdf" and not raw_text.strip():
                    image_paths = _pdf_to_images(file_path)
                    for img_path in image_paths:
                        img_rel = str(img_path.relative_to(case_dir))
                        manifest.files.append(FileEntry(
                            path=img_rel,
                            size=img_path.stat().st_size,
                            mime_type="image/png",
                            preview=_extract_text_image(img_path),
                            checksum=_checksum(img_path),
                        ))
                    # Skip the original empty PDF entry
                    continue

                preview = raw_text[:PREVIEW_CHARS]
                csum = _checksum(file_path)

                manifest.files.append(
                    FileEntry(
                        path=rel_path,
                        size=size,
                        mime_type=mime,
                        preview=preview,
                        checksum=csum,
                    )
                )

        manifest.total_files = len(manifest.files)
        return manifest
